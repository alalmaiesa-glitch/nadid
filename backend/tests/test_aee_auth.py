from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.main import app


def make_docx() -> bytes:
    document = Document()
    document.add_heading("اختبار نَضِيد", level=1)
    document.add_paragraph(
        "يساهم ذلك في تحسين من مستوى الأداء، وتبلغ التكلفة 100 ريال."
    )
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_health_is_public(monkeypatch):
    monkeypatch.setenv("NADID_ENV", "production")
    monkeypatch.setenv("AEE_INTERNAL_TOKEN", "test-secret")

    client = TestClient(app)
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_production_aee_rejects_missing_token(monkeypatch):
    monkeypatch.setenv("NADID_ENV", "production")
    monkeypatch.setenv("AEE_INTERNAL_TOKEN", "test-secret")

    client = TestClient(app)
    response = client.post(
        "/v1/analyze/docx",
        files={
            "file": (
                "sample.docx",
                make_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 401


def test_production_aee_rejects_wrong_token(monkeypatch):
    monkeypatch.setenv("NADID_ENV", "production")
    monkeypatch.setenv("AEE_INTERNAL_TOKEN", "test-secret")

    client = TestClient(app)
    response = client.post(
        "/v1/analyze/docx",
        headers={"Authorization": "Bearer wrong-secret"},
        files={
            "file": (
                "sample.docx",
                make_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 401


def test_production_aee_accepts_internal_token(monkeypatch):
    monkeypatch.setenv("NADID_ENV", "production")
    monkeypatch.setenv("AEE_INTERNAL_TOKEN", "test-secret")

    client = TestClient(app)
    response = client.post(
        "/v1/analyze/docx",
        headers={"Authorization": "Bearer test-secret"},
        files={
            "file": (
                "sample.docx",
                make_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["document"]["word_count"] > 0
