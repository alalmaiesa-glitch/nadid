from __future__ import annotations

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from fastapi.testclient import TestClient

from app.main import app


client = TestClient(app, raise_server_exceptions=False)


def _valid_docx(text: str = "هاذا نص تجريبي للمراجعة.") -> bytes:
    document = Document()
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _rewrite_package(
    payload: bytes,
    replacements: dict[str, bytes] | None = None,
    removals: set[str] | None = None,
) -> bytes:
    replacements = replacements or {}
    removals = removals or set()
    output = BytesIO()

    with ZipFile(BytesIO(payload), "r") as source:
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
            for info in source.infolist():
                if info.filename in removals:
                    continue
                data = replacements.get(
                    info.filename,
                    source.read(info.filename),
                )
                target.writestr(info.filename, data)

    return output.getvalue()


def _post(path: str, payload: bytes, filename: str = "test.docx"):
    data = {}
    if path == "/v1/apply/docx":
        data["patches"] = "[]"
    return client.post(
        path,
        files={
            "file": (
                filename,
                payload,
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            )
        },
        data=data,
    )


def test_corrupt_document_xml_is_422_for_all_docx_endpoints():
    payload = _rewrite_package(
        _valid_docx(),
        {"word/document.xml": b"<w:document"},
    )

    for path in (
        "/v1/analyze/docx",
        "/v1/analyze/docx/deep",
        "/v1/apply/docx",
    ):
        response = _post(path, payload)
        assert response.status_code == 422, (path, response.text)
        assert "invalid_xml_part" in response.text


def test_missing_package_relationships_fails_closed():
    payload = _rewrite_package(
        _valid_docx(),
        removals={"_rels/.rels"},
    )

    for path in (
        "/v1/analyze/docx",
        "/v1/analyze/docx/deep",
        "/v1/apply/docx",
    ):
        response = _post(path, payload)
        assert response.status_code == 422, (path, response.text)


def test_blank_document_has_no_reviewable_content():
    document = Document()
    stream = BytesIO()
    document.save(stream)

    for path in ("/v1/analyze/docx", "/v1/analyze/docx/deep"):
        response = _post(path, stream.getvalue())
        assert response.status_code == 422
        assert "No reviewable content found" in response.text
