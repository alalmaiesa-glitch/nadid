from app.contracts import DocumentNode, ProtectedSpan
from app.pipeline.chunker import build_chunks
from app.pipeline.patch_validator import validate_patch
from app.pipeline.reviewer import fast_review


def node(text: str, sequence: int = 0) -> DocumentNode:
    return DocumentNode(
        id=f"n-{sequence}",
        type="paragraph",
        text=text,
        sequence_no=sequence,
    )


def test_fast_review_detects_extra_preposition():
    issues = fast_review([node("يساهم ذلك في تحسين من مستوى الأداء.")])
    assert any(issue.replacement == "تحسين مستوى" for issue in issues)


def test_chunker_preserves_nodes():
    nodes = [
        node("فقرة قصيرة.", 0),
        node("فقرة أخرى.", 1),
    ]
    chunks = build_chunks(nodes, target_tokens=100)
    assert len(chunks) == 1
    assert chunks[0].node_ids == ["n-0", "n-1"]


def test_fact_lock_blocks_number_change():
    text = "بلغت التكلفة 38,771,251 ريال."
    protected = [
        ProtectedSpan(
            id="p1",
            node_id="n-0",
            type="currency",
            value="38,771,251 ريال",
            validator_key="numeric_equivalence",
        )
    ]

    result = validate_patch(
        block_text=text,
        original="38,771,251",
        replacement="39,771,251",
        protected_spans=protected,
    )

    assert result.status == "BLOCK"
    assert result.candidate == text


def test_fact_lock_allows_safe_style_edit():
    text = "بلغت التكلفة 38,771,251 ريال ، وهي معتمدة."
    protected = [
        ProtectedSpan(
            id="p1",
            node_id="n-0",
            type="currency",
            value="38,771,251 ريال",
            validator_key="numeric_equivalence",
        )
    ]

    result = validate_patch(
        block_text=text,
        original="ريال ،",
        replacement="ريال،",
        protected_spans=protected,
    )

    assert result.status == "PASS"
    assert "38,771,251 ريال" in result.candidate


def test_fact_lock_blocks_negation_removal():
    text = "ولا يسمح النظام بتغيير القيمة."
    protected = [
        ProtectedSpan(
            id="p2",
            node_id="n-0",
            type="negation",
            value="ولا يسمح النظام بتغيير القيمة",
            validator_key="negation_preservation",
            lock_policy="semantic_exact",
        )
    ]

    result = validate_patch(
        block_text=text,
        original="ولا يسمح",
        replacement="ويسمح",
        protected_spans=protected,
    )

    assert result.status == "BLOCK"


def test_document_memory_extracts_repeated_terms():
    from app.pipeline.chunker import build_chunks
    from app.pipeline.memory import build_document_memory
    from app.pipeline.protection import extract_protected_spans

    nodes = [
        node("يعتمد المشروع على الذكاء الاصطناعي في التحليل.", 0),
        node("يستخدم الذكاء الاصطناعي لتحسين المراجعة.", 1),
    ]
    chunks = build_chunks(nodes)
    protected = extract_protected_spans(nodes)
    memory = build_document_memory(nodes, chunks, protected)

    assert any(term.term == "الذكاء" and term.count >= 2 for term in memory.terms)


def test_fact_conflict_detects_same_claim_different_value():
    from app.pipeline.facts import detect_fact_conflicts, extract_facts

    nodes = [
        node("عدد المحاور في النموذج 4 محاور.", 0),
        node("عدد المحاور في النموذج 5 محاور.", 1),
    ]
    facts = extract_facts(nodes)
    conflicts = detect_fact_conflicts(facts)

    assert len(conflicts) == 1
    assert set(conflicts[0].values) == {"4", "5"}


def test_context_retrieval_finds_related_chunk():
    from app.pipeline.chunker import build_chunks
    from app.pipeline.context import retrieve_context
    from app.pipeline.memory import build_document_memory
    from app.pipeline.protection import extract_protected_spans

    nodes = [
        node("يعالج الفصل الأول الجوانب التشغيلية العامة.", 0),
        node("تبلغ الطاقة الإنتاجية للمصنع 2000000 وحدة سنويًا.", 1),
        node("يتناول هذا الجزء استراتيجية التسويق والتوزيع.", 2),
        node("تساعد الطاقة الإنتاجية في تقدير احتياجات المواد.", 3),
    ]
    chunks = build_chunks(nodes, target_tokens=7, hard_limit=15)
    protected = extract_protected_spans(nodes)
    memory = build_document_memory(nodes, chunks, protected)

    package = retrieve_context(
        target_node_id="n-1",
        nodes=nodes,
        chunks=chunks,
        memory=memory,
        query="الطاقة الإنتاجية",
        max_chunks=3,
    )

    assert any("الطاقة الإنتاجية" in hit.text for hit in package.hits)


def test_parser_node_ids_are_stable_for_same_docx():
    from io import BytesIO
    from docx import Document
    from app.pipeline.parser import parse_docx

    document = Document()
    document.add_heading("مقدمة", level=1)
    document.add_paragraph("هذا نص تجريبي ثابت.")
    stream = BytesIO()
    document.save(stream)
    payload = stream.getvalue()

    first = parse_docx(payload)
    second = parse_docx(payload)

    assert [item.id for item in first] == [item.id for item in second]


def test_chunk_ids_are_stable():
    nodes = [
        node("الفقرة الأولى في المستند.", 0),
        node("الفقرة الثانية في المستند.", 1),
    ]

    first = build_chunks(nodes, target_tokens=100)
    second = build_chunks(nodes, target_tokens=100)

    assert [item.id for item in first] == [item.id for item in second]


def test_fact_ids_are_stable():
    from app.pipeline.facts import extract_facts

    nodes = [node("تبلغ الطاقة الإنتاجية 2000000 وحدة سنويًا.", 0)]

    first = extract_facts(nodes)
    second = extract_facts(nodes)

    assert [item.id for item in first] == [item.id for item in second]


def test_arabic_numeric_forms_are_canonicalized():
    from app.arabic_numbers import canonical_decimal

    assert canonical_decimal("٣٨٬٧٧١٬٢٥١ ريال") == "38771251"
    assert canonical_decimal("١٨٫٦٪") == "18.6"


def test_protection_extracts_arabic_currency_percentage_and_year():
    from app.pipeline.protection import extract_protected_spans

    spans = extract_protected_spans([
        node("بلغت التكلفة ٣٨٬٧٧١٬٢٥١ ريال في عام ٢٠٢٦ بنسبة ١٨٫٦٪.", 0)
    ])

    values = {(span.type, span.value) for span in spans}

    assert ("currency", "٣٨٬٧٧١٬٢٥١ ريال") in values
    assert ("date", "٢٠٢٦") in values
    assert ("percentage", "١٨٫٦٪") in values


def test_fact_lock_allows_equivalent_numeric_format():
    text = "بلغت التكلفة ٣٨٬٧٧١٬٢٥١ ريال."
    protected = [
        ProtectedSpan(
            id="arabic-number",
            node_id="n-0",
            type="currency",
            value="٣٨٬٧٧١٬٢٥١ ريال",
            validator_key="numeric_equivalence",
        )
    ]

    result = validate_patch(
        block_text=text,
        original="٣٨٬٧٧١٬٢٥١ ريال",
        replacement="38,771,251 ريال",
        protected_spans=protected,
    )

    assert result.status == "PASS"


def test_fact_lock_blocks_changed_arabic_number():
    text = "بلغت التكلفة ٣٨٬٧٧١٬٢٥١ ريال."
    protected = [
        ProtectedSpan(
            id="arabic-number",
            node_id="n-0",
            type="currency",
            value="٣٨٬٧٧١٬٢٥١ ريال",
            validator_key="numeric_equivalence",
        )
    ]

    result = validate_patch(
        block_text=text,
        original="٣٨٬٧٧١٬٢٥١",
        replacement="٣٩٬٧٧١٬٢٥١",
        protected_spans=protected,
    )

    assert result.status == "BLOCK"


def test_docx_patch_preserves_surrounding_run_formatting():
    from io import BytesIO
    from docx import Document
    from app.pipeline.docx_patch import apply_patches_to_docx
    from app.pipeline.parser import parse_docx
    from app.contracts import PatchOperation

    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("بلغت التكلفة ")
    first.bold = True
    paragraph.add_run("38,771,251 ريال ، وهي معتمدة.")

    source = BytesIO()
    document.save(source)
    payload = source.getvalue()

    nodes = parse_docx(payload)
    target = nodes[0]

    output, report = apply_patches_to_docx(
        payload,
        [
            PatchOperation(
                node_id=target.id,
                original="ريال ،",
                replacement="ريال،",
            )
        ],
    )

    assert len(report.applied) == 1
    assert not report.skipped

    result = Document(BytesIO(output))
    result_paragraph = result.paragraphs[0]

    assert "38,771,251 ريال، وهي معتمدة." in result_paragraph.text
    assert result_paragraph.runs[0].bold is True
