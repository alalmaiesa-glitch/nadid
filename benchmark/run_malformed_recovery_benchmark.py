from __future__ import annotations

import argparse
import json
import time
import warnings
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from docx import Document
from fastapi.testclient import TestClient

from app.docx_security import (
    MAX_WORD_COUNT,
    UnsafeDocxError,
    validate_word_count,
)
from app.main import app


ROOT = Path(__file__).resolve().parent
DEFAULT_CONFIG = ROOT / "malformed" / "v1.json"
CLIENT = TestClient(app, raise_server_exceptions=False)


def _docx_with_paragraph(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _blank_docx() -> bytes:
    document = Document()
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _dense_table_docx() -> bytes:
    document = Document()
    table = document.add_table(rows=30, cols=30)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = (
                f"هاذا البند {row_index + 1}-{col_index + 1} "
                "ضمن الجدول التشغيلي."
            )
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _rewrite_package(
    payload: bytes,
    replacements: dict[str, bytes] | None = None,
    removals: set[str] | None = None,
) -> bytes:
    replacements = replacements or {}
    removals = removals or set()
    output = BytesIO()

    with ZipFile(BytesIO(payload), "r") as source:
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
            for info in source.infolist():
                if info.filename in removals:
                    continue
                data = replacements.get(
                    info.filename,
                    source.read(info.filename),
                )
                target.writestr(info.filename, data)

    return output.getvalue()


def _append_entry(payload: bytes, name: str, value: bytes) -> bytes:
    output = BytesIO(payload)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with ZipFile(output, "a", compression=ZIP_DEFLATED) as archive:
            archive.writestr(name, value)
    return output.getvalue()


def _minimal_zip(
    entries: dict[str, bytes],
    compression: int = ZIP_DEFLATED,
) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", compression=compression) as archive:
        for name, value in entries.items():
            archive.writestr(name, value)
    return output.getvalue()


def _too_many_entries() -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_STORED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("word/document.xml", b"<document/>")
        for index in range(20_001):
            archive.writestr(f"extra/e-{index}.txt", b"")
    return output.getvalue()


def _fixture(kind: str, valid: bytes) -> bytes:
    if kind == "valid":
        return valid
    if kind == "empty":
        return b""
    if kind == "plain_bytes":
        return b"this is not a docx package"
    if kind == "truncated_zip":
        return valid[: max(64, len(valid) // 3)]
    if kind == "missing_content_types":
        return _rewrite_package(
            valid,
            removals={"[Content_Types].xml"},
        )
    if kind == "missing_document_xml":
        return _rewrite_package(
            valid,
            removals={"word/document.xml"},
        )
    if kind == "zip_traversal":
        return _append_entry(valid, "../escape.txt", b"x")
    if kind == "zip_absolute_path":
        return _append_entry(valid, "/absolute.txt", b"x")
    if kind == "zip_backslash_traversal":
        return _append_entry(valid, "..\\escape.txt", b"x")
    if kind == "duplicate_document_entry":
        return _append_entry(
            valid,
            "word/document.xml",
            b"<duplicate/>",
        )
    if kind == "corrupt_document_xml":
        return _rewrite_package(
            valid,
            replacements={"word/document.xml": b"<w:document"},
        )
    if kind == "missing_package_rels":
        return _rewrite_package(
            valid,
            removals={"_rels/.rels"},
        )
    if kind == "blank_document":
        return _blank_docx()
    if kind == "whitespace_document":
        return _docx_with_paragraph("   \t   ")
    if kind == "huge_single_paragraph":
        return _docx_with_paragraph(
            ("هاذا النص الطويل للاختبار " * 20_000).strip()
        )
    if kind == "dense_table":
        return _dense_table_docx()
    if kind == "too_many_entries":
        return _too_many_entries()
    if kind == "invalid_content_types_xml":
        return _rewrite_package(
            valid,
            replacements={"[Content_Types].xml": b"<Types"},
        )
    if kind == "fake_docx_package":
        return _minimal_zip(
            {
                "[Content_Types].xml": b"<Types/>",
                "word/document.xml": b"<document/>",
            }
        )
    raise KeyError(kind)


def _post(endpoint: str, filename: str, payload: bytes):
    path = {
        "analyze": "/v1/analyze/docx",
        "deep": "/v1/analyze/docx/deep",
        "apply": "/v1/apply/docx",
    }[endpoint]
    form = {"patches": "[]"} if endpoint == "apply" else {}
    return CLIENT.post(
        path,
        files={
            "file": (
                filename,
                payload,
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            )
        },
        data=form,
    )


def _detail_text(response) -> str:
    try:
        body = response.json()
    except Exception:
        return response.text
    detail = body.get("detail") if isinstance(body, dict) else body
    return json.dumps(detail, ensure_ascii=False, sort_keys=True)


def _run_case(case: dict[str, Any], valid: bytes) -> dict[str, Any]:
    started = time.perf_counter()
    failures: list[str] = []
    observed: dict[str, Any] = {}

    if case["fixture"] == "word_count_limit":
        try:
            validate_word_count(MAX_WORD_COUNT + 1)
        except UnsafeDocxError as exc:
            observed["status"] = "blocked"
            observed["detail"] = str(exc)
        else:
            observed["status"] = "allowed"
            observed["detail"] = ""

        if observed["status"] != case["expected_status"]:
            failures.append(
                f"status:{observed['status']}!={case['expected_status']}"
            )
        expected_detail = case.get("detail_contains")
        if expected_detail and expected_detail not in observed["detail"]:
            failures.append("missing_expected_detail")
    else:
        payload = _fixture(case["fixture"], valid)
        response = _post(
            case["endpoint"],
            case.get("filename", "document.docx"),
            payload,
        )
        detail = (
            ""
            if response.status_code == 200 and case["endpoint"] == "apply"
            else _detail_text(response)
        )
        observed = {
            "status": response.status_code,
            "detail": detail[:500],
            "payload_bytes": len(payload),
        }

        if response.status_code != case["expected_status"]:
            failures.append(
                f"status:{response.status_code}!={case['expected_status']}"
            )

        expected_detail = case.get("detail_contains")
        if expected_detail and expected_detail not in detail:
            failures.append("missing_expected_detail")

        if response.status_code >= 500:
            failures.append("server_error_escape")

        if response.status_code == 200 and case.get("min_nodes"):
            body = response.json()
            if case["endpoint"] == "deep":
                nodes = body["base"]["nodes"]
            else:
                nodes = body["nodes"]
            observed["node_count"] = len(nodes)
            if len(nodes) < int(case["min_nodes"]):
                failures.append(
                    f"node_count:{len(nodes)}<{case['min_nodes']}"
                )

    return {
        "id": case["id"],
        "fixture": case["fixture"],
        "endpoint": case["endpoint"],
        "expected_status": case["expected_status"],
        "observed": observed,
        "elapsed_seconds": round(time.perf_counter() - started, 6),
        "failures": failures,
        "passed": not failures,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--enforce", action="store_true")
    args = parser.parse_args()

    config = json.loads(args.config.read_text(encoding="utf-8"))
    valid = _docx_with_paragraph("هاذا التقرير يحتاج إلى مراجعة.")
    results = [_run_case(case, valid) for case in config["cases"]]

    failed = [
        result["id"]
        for result in results
        if not result["passed"]
    ]
    rejected = sum(
        1
        for result in results
        if isinstance(result["observed"].get("status"), int)
        and result["observed"]["status"] >= 400
    )
    accepted = sum(
        1
        for result in results
        if result["observed"].get("status") == 200
    )

    report = {
        "schema_version": 1,
        "benchmark": "malformed_document_failure_recovery_v1",
        "cases": results,
        "summary": {
            "passed": not failed,
            "passed_cases": len(results) - len(failed),
            "total_cases": len(results),
            "failed_cases": failed,
            "rejected_http_cases": rejected,
            "accepted_http_cases": accepted,
            "server_errors": sum(
                1
                for result in results
                if result["observed"].get("status", 0) in range(500, 600)
            ),
        },
    }

    rendered = json.dumps(
        report,
        ensure_ascii=False,
        indent=2,
        sort_keys=True,
    )
    print(rendered)

    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(rendered + "\n", encoding="utf-8")

    if args.enforce and failed:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
