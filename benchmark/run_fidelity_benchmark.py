from __future__ import annotations

import argparse
import json
import struct
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

from app.contracts import PatchOperation
from app.pipeline.auto_apply import build_safe_auto_apply_plan
from app.pipeline.docx_fidelity import verify_docx_fidelity
from app.pipeline.docx_patch import apply_patches_to_docx
from app.pipeline.parser import parse_docx
from app.pipeline.protection import extract_protected_spans
from app.pipeline.reviewer import fast_review


def _save(document: Document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _bmp_1x1() -> bytes:
    file_header = struct.pack(
        "<2sIHHI",
        b"BM",
        58,
        0,
        0,
        54,
    )
    dib = struct.pack(
        "<IIIHHIIIIII",
        40,
        1,
        1,
        1,
        24,
        0,
        4,
        2835,
        2835,
        0,
        0,
    )
    pixel = bytes([0, 90, 220, 0])
    return file_header + dib + pixel


def _add_hyperlink(paragraph, text: str, url: str):
    relation_id = paragraph.part.relate_to(
        url,
        RT.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)

    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)

    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return relation_id


def _add_bookmark(paragraph, name: str = "nadid-bookmark"):
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), "42")
    bookmark_start.set(qn("w:name"), name)
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), "42")

    paragraph._p.insert(0, bookmark_start)
    paragraph._p.append(bookmark_end)


def _auto_apply(data: bytes):
    nodes = parse_docx(data)
    protected = extract_protected_spans(nodes)
    suggestions = fast_review(nodes)
    plan = build_safe_auto_apply_plan(
        nodes,
        suggestions,
        protected,
    )
    output, report = apply_patches_to_docx(
        data,
        plan.patches,
    )
    return output, report, plan


def _media_hashes(data: bytes) -> dict[str, bytes]:
    with ZipFile(BytesIO(data), "r") as archive:
        return {
            name: archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/media/")
        }


def _relationship_targets(data: bytes) -> set[str]:
    with ZipFile(BytesIO(data), "r") as archive:
        payload = archive.read("word/_rels/document.xml.rels")
    text = payload.decode("utf-8")
    targets = set()
    marker = 'Target="'
    cursor = 0
    while True:
        index = text.find(marker, cursor)
        if index < 0:
            break
        start = index + len(marker)
        end = text.find('"', start)
        if end < 0:
            break
        targets.add(text[start:end])
        cursor = end + 1
    return targets


def _document_xml(data: bytes) -> str:
    with ZipFile(BytesIO(data), "r") as archive:
        return archive.read("word/document.xml").decode("utf-8")


def _fixture(scenario: str) -> bytes:
    document = Document()

    if scenario == "plain_paragraph":
        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "bold_surrounding":
        paragraph = document.add_paragraph()
        bold = paragraph.add_run("بيان مهم: ")
        bold.bold = True
        paragraph.add_run("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "italic_underline_font":
        paragraph = document.add_paragraph()
        run = paragraph.add_run("مرجع تمهيدي. ")
        run.italic = True
        run.underline = True
        run.font.name = "Arial"
        run.font.size = Pt(12)
        paragraph.add_run("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "heading_style":
        document.add_heading("هاذا العنوان", level=2)
        document.add_paragraph("المتن سليم.")
        return _save(document)

    if scenario == "table_cell":
        document.add_paragraph("مقدمة سليمة.")
        table = document.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "البند"
        table.cell(0, 1).text = "القيمة"
        table.cell(1, 0).text = "هاذا البيان"
        table.cell(1, 1).text = "100"
        return _save(document)

    if scenario == "merged_table":
        table = document.add_table(rows=2, cols=3)
        table.style = "Table Grid"
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "عنوان مدمج"
        table.cell(0, 2).text = "ثابت"
        table.cell(1, 0).text = "هاذا البيان"
        table.cell(1, 1).text = "قيمة"
        table.cell(1, 2).text = "ملاحظة"
        return _save(document)

    if scenario == "numbered_list":
        document.add_paragraph("البند الأول", style="List Number")
        document.add_paragraph("البند الثاني", style="List Number")
        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "bullet_list":
        document.add_paragraph("النقطة الأولى", style="List Bullet")
        document.add_paragraph("النقطة الثانية", style="List Bullet")
        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "image_relationship":
        paragraph = document.add_paragraph("صورة توضيحية: ")
        run = paragraph.add_run()
        run.add_picture(BytesIO(_bmp_1x1()), width=Inches(0.25))
        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "hyperlink_relationship":
        paragraph = document.add_paragraph("للمزيد: ")
        _add_hyperlink(
            paragraph,
            "الموقع المرجعي",
            "https://example.com/nadid-fidelity",
        )
        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "header_footer":
        section = document.sections[0]
        section.header.paragraphs[0].text = "رأس مستند ثابت"
        section.footer.paragraphs[0].text = "تذييل مستند ثابت"
        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "page_margins_orientation":
        section = document.sections[0]
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(1.0)
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = (
            section.page_height,
            section.page_width,
        )
        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "multiple_sections":
        document.add_paragraph("القسم الأول ثابت.")
        second = document.add_section(WD_SECTION_START.NEW_PAGE)
        second.top_margin = Inches(0.6)
        second.header.paragraphs[0].text = "رأس القسم الثاني"
        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "tab_and_break":
        paragraph = document.add_paragraph()
        paragraph.add_run("هاذا التقرير")
        special = paragraph.add_run()
        special.add_tab()
        special.add_text("تفصيل ثابت")
        special.add_break(WD_BREAK.LINE)
        special.add_text("سطر ثابت")
        return _save(document)

    if scenario == "bookmark":
        paragraph = document.add_paragraph("موضع إشارة ثابت")
        _add_bookmark(paragraph)
        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "leading_whitespace":
        document.add_paragraph("   هاذا التقرير معتمد.   ")
        return _save(document)

    if scenario == "repeated_exact_spans":
        document.add_paragraph(
            "هاذا التقرير الأول، وهاذا التقرير الثاني."
        )
        return _save(document)

    if scenario == "cross_run_same_format":
        paragraph = document.add_paragraph()
        first = paragraph.add_run("ها")
        second = paragraph.add_run("ذا")
        first.bold = True
        second.bold = True
        paragraph.add_run(" التقرير معتمد.")
        return _save(document)

    if scenario == "cross_run_mixed_format":
        paragraph = document.add_paragraph()
        first = paragraph.add_run("ها")
        second = paragraph.add_run("ذا")
        first.bold = True
        second.italic = True
        paragraph.add_run(" التقرير معتمد.")
        return _save(document)

    if scenario == "unsafe_special_run_fails_closed":
        paragraph = document.add_paragraph()
        paragraph.add_run("ه")
        special = paragraph.add_run()
        special.add_tab()
        paragraph.add_run("اذا")
        return _save(document)

    if scenario in {
        "no_patch_byte_identity",
        "second_round_byte_identity",
    }:
        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "combined_rich_document":
        section = document.sections[0]
        section.header.paragraphs[0].text = "رأس غني"
        section.footer.paragraphs[0].text = "تذييل غني"
        section.top_margin = Inches(0.65)

        document.add_heading("عنوان ثابت", level=1)
        document.add_paragraph("بند ثابت", style="List Number")
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "خلية ثابتة"
        table.cell(0, 1).text = "قيمة ثابتة"

        image_paragraph = document.add_paragraph("صورة: ")
        image_paragraph.add_run().add_picture(
            BytesIO(_bmp_1x1()),
            width=Inches(0.2),
        )

        link_paragraph = document.add_paragraph("رابط: ")
        _add_hyperlink(
            link_paragraph,
            "مرجع ثابت",
            "https://example.com/rich",
        )

        document.add_paragraph("هاذا التقرير معتمد.")
        return _save(document)

    if scenario == "target_run_format_preserved":
        paragraph = document.add_paragraph()
        target = paragraph.add_run("هاذا")
        target.bold = True
        target.underline = True
        target.font.name = "Arial"
        target.font.size = Pt(14)
        paragraph.add_run(" التقرير معتمد.")
        return _save(document)

    raise ValueError(f"Unknown scenario: {scenario}")


def _direct_special_run_patch(data: bytes):
    nodes = parse_docx(data)
    target = nodes[0]
    patch = PatchOperation(
        node_id=target.id,
        original="ه\tاذا",
        replacement="هذا",
        start_offset=0,
    )
    output, report = apply_patches_to_docx(
        data,
        [patch],
    )
    return output, report, [patch]


def _scenario_assertions(
    scenario: str,
    before: bytes,
    after: bytes,
    report,
    plan,
) -> list[str]:
    failures: list[str] = []

    if not report.fidelity_ok:
        failures.extend(
            f"fidelity:{error}"
            for error in (report.fidelity_errors or [])
        )

    if scenario not in {
        "unsafe_special_run_fails_closed",
        "no_patch_byte_identity",
    }:
        if not report.applied:
            failures.append("no_patch_applied")

    before_doc = Document(BytesIO(before))
    after_doc = Document(BytesIO(after))

    if scenario == "bold_surrounding":
        if after_doc.paragraphs[0].runs[0].bold is not True:
            failures.append("bold_surrounding_lost")

    if scenario == "italic_underline_font":
        run = after_doc.paragraphs[0].runs[0]
        if run.italic is not True:
            failures.append("italic_lost")
        if run.underline is not True:
            failures.append("underline_lost")
        if run.font.name != "Arial":
            failures.append("font_name_lost")
        if run.font.size != Pt(12):
            failures.append("font_size_lost")

    if scenario == "heading_style":
        if after_doc.paragraphs[0].style.name != before_doc.paragraphs[0].style.name:
            failures.append("heading_style_changed")

    if scenario in {"table_cell", "merged_table"}:
        if len(after_doc.tables) != len(before_doc.tables):
            failures.append("table_count_changed")
        if after_doc.tables[0].style.name != before_doc.tables[0].style.name:
            failures.append("table_style_changed")

    if scenario == "merged_table":
        before_grid = _document_xml(before).count("<w:gridSpan")
        after_grid = _document_xml(after).count("<w:gridSpan")
        if before_grid != after_grid:
            failures.append("merged_cell_geometry_changed")

    if scenario in {"numbered_list", "bullet_list"}:
        before_styles = [
            paragraph.style.name
            for paragraph in before_doc.paragraphs[:2]
        ]
        after_styles = [
            paragraph.style.name
            for paragraph in after_doc.paragraphs[:2]
        ]
        if before_styles != after_styles:
            failures.append("list_styles_changed")

    if scenario in {"image_relationship", "combined_rich_document"}:
        if _media_hashes(before) != _media_hashes(after):
            failures.append("media_payload_changed")
        if not _media_hashes(after):
            failures.append("media_missing")

    if scenario in {"hyperlink_relationship", "combined_rich_document"}:
        before_targets = _relationship_targets(before)
        after_targets = _relationship_targets(after)
        if before_targets != after_targets:
            failures.append("relationship_targets_changed")
        if not any(
            target.startswith("https://example.com/")
            for target in after_targets
        ):
            failures.append("hyperlink_target_missing")

    if scenario in {"header_footer", "combined_rich_document"}:
        if (
            before_doc.sections[0].header.paragraphs[0].text
            != after_doc.sections[0].header.paragraphs[0].text
        ):
            failures.append("header_changed")
        if (
            before_doc.sections[0].footer.paragraphs[0].text
            != after_doc.sections[0].footer.paragraphs[0].text
        ):
            failures.append("footer_changed")

    if scenario == "page_margins_orientation":
        before_section = before_doc.sections[0]
        after_section = after_doc.sections[0]
        attributes = [
            "top_margin",
            "bottom_margin",
            "left_margin",
            "right_margin",
            "page_width",
            "page_height",
            "orientation",
        ]
        for attribute in attributes:
            if (
                getattr(before_section, attribute)
                != getattr(after_section, attribute)
            ):
                failures.append(
                    f"section_property_changed:{attribute}"
                )

    if scenario == "multiple_sections":
        if len(before_doc.sections) != len(after_doc.sections):
            failures.append("section_count_changed")
        if (
            before_doc.sections[1].top_margin
            != after_doc.sections[1].top_margin
        ):
            failures.append("second_section_margin_changed")

    if scenario == "tab_and_break":
        before_xml = _document_xml(before)
        after_xml = _document_xml(after)
        if before_xml.count("<w:tab") != after_xml.count("<w:tab"):
            failures.append("tab_changed")
        if before_xml.count("<w:br") != after_xml.count("<w:br"):
            failures.append("break_changed")

    if scenario == "bookmark":
        before_xml = _document_xml(before)
        after_xml = _document_xml(after)
        for tag in ("bookmarkStart", "bookmarkEnd"):
            if before_xml.count(f"<w:{tag}") != after_xml.count(f"<w:{tag}"):
                failures.append(f"{tag}_changed")

    if scenario == "leading_whitespace":
        if not after_doc.paragraphs[0].text.startswith("   هذا"):
            failures.append("leading_whitespace_or_offset_changed")

    if scenario == "repeated_exact_spans":
        text = after_doc.paragraphs[0].text
        if text.count("هذا") != 2 or "هاذا" in text:
            failures.append("repeated_spans_not_all_fixed")

    if scenario == "cross_run_same_format":
        runs = after_doc.paragraphs[0].runs
        if runs[0].bold is not True or runs[1].bold is not True:
            failures.append("same_format_run_properties_changed")
        if "".join(run.text for run in runs[:2]) != "هذا":
            failures.append("cross_run_text_wrong")

    if scenario == "cross_run_mixed_format":
        runs = after_doc.paragraphs[0].runs
        if runs[0].bold is not True:
            failures.append("mixed_run_bold_lost")
        if runs[1].italic is not True:
            failures.append("mixed_run_italic_lost")
        if "".join(run.text for run in runs[:2]) != "هذا":
            failures.append("mixed_run_text_wrong")

    if scenario == "unsafe_special_run_fails_closed":
        if report.applied:
            failures.append("unsafe_special_run_applied")
        if not report.skipped:
            failures.append("unsafe_special_run_not_skipped")
        if after != before:
            failures.append("unsafe_skip_changed_bytes")

    if scenario == "no_patch_byte_identity":
        if after != before:
            failures.append("no_patch_not_byte_identical")

    if scenario == "second_round_byte_identity":
        second_nodes = parse_docx(after)
        second_plan = build_safe_auto_apply_plan(
            second_nodes,
            fast_review(second_nodes),
            extract_protected_spans(second_nodes),
        )
        second_output, second_report = apply_patches_to_docx(
            after,
            second_plan.patches,
        )
        if second_plan.patches:
            failures.append("second_round_still_has_auto_fixes")
        if second_output != after:
            failures.append("second_round_not_byte_identical")
        if not second_report.fidelity_ok:
            failures.append("second_round_fidelity_failed")

    if scenario == "target_run_format_preserved":
        run = after_doc.paragraphs[0].runs[0]
        if run.bold is not True:
            failures.append("target_bold_lost")
        if run.underline is not True:
            failures.append("target_underline_lost")
        if run.font.name != "Arial":
            failures.append("target_font_lost")
        if run.font.size != Pt(14):
            failures.append("target_font_size_lost")

    return failures


def evaluate(case: dict) -> dict:
    scenario = case["scenario"]
    before = _fixture(scenario)

    if scenario == "unsafe_special_run_fails_closed":
        after, report, patches = _direct_special_run_patch(before)
        plan = type("DirectPlan", (), {"patches": patches})()
    elif scenario == "no_patch_byte_identity":
        after, report = apply_patches_to_docx(before, [])
        plan = type("NoPlan", (), {"patches": []})()
    else:
        after, report, plan = _auto_apply(before)

    fidelity = verify_docx_fidelity(
        before,
        after,
        [
            patch
            for patch in plan.patches
            if patch.node_id in report.applied
        ],
    )

    failures = _scenario_assertions(
        scenario,
        before,
        after,
        report,
        plan,
    )

    if report.applied and not fidelity.ok:
        failures.extend(
            f"external_fidelity:{error}"
            for error in fidelity.errors
        )

    return {
        "id": case["id"],
        "scenario": scenario,
        "passed": not failures,
        "failures": failures,
        "applied": len(report.applied),
        "skipped": len(report.skipped),
        "fidelity_ok": report.fidelity_ok,
        "checks": fidelity.checks,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--suite",
        default="benchmark/fidelity/v1.json",
    )
    parser.add_argument("--report")
    parser.add_argument("--enforce", action="store_true")
    args = parser.parse_args()

    payload = json.loads(
        Path(args.suite).read_text(encoding="utf-8")
    )
    results = [evaluate(case) for case in payload["cases"]]
    passed = sum(item["passed"] for item in results)
    total = len(results)
    score = passed / total if total else 0.0

    report = {
        "schema_version": 1,
        "suite": args.suite,
        "passed": passed,
        "total": total,
        "score": round(score, 4),
        "failures": [
            item for item in results if not item["passed"]
        ],
        "results": results,
    }

    print(
        f"Document Fidelity & Round-Trip Safety V1: "
        f"{passed}/{total} = {score:.1%}"
    )
    if report["failures"]:
        print("\nFailures:")
        for item in report["failures"]:
            print(
                f"- {item['id']}: "
                + ", ".join(item["failures"])
            )

    if args.report:
        Path(args.report).write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    if args.enforce and report["failures"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
