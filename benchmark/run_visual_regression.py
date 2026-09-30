from __future__ import annotations

import argparse
import json
import math
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

import fitz
from PIL import Image, ImageChops
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

from app.pipeline.auto_apply import build_safe_auto_apply_plan
from app.pipeline.docx_patch import apply_patches_to_docx
from app.pipeline.parser import parse_docx
from app.pipeline.protection import extract_protected_spans
from app.pipeline.reviewer import fast_review


def _save(document: Document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _set_rtl(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")


def _add_anchor(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.font.name = "DejaVu Sans"
    run.font.size = Pt(8)
    return paragraph


def _add_hyperlink(paragraph, text: str, url: str):
    relation_id = paragraph.part.relate_to(
        url,
        RT.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _bmp_1x1() -> bytes:
    # Minimal 1x1 24-bit BMP.
    return (
        b"BM:\x00\x00\x00\x00\x00\x00\x006\x00\x00\x00"
        b"(\x00\x00\x00\x01\x00\x00\x00\x01\x00\x00\x00"
        b"\x01\x00\x18\x00\x00\x00\x00\x00\x04\x00\x00\x00"
        b"\x13\x0b\x00\x00\x13\x0b\x00\x00\x00\x00\x00\x00"
        b"\x00\x00\x00\x00\x00\x90\xdc\x00"
    )


def _arabic_paragraph(document: Document, text: str):
    paragraph = document.add_paragraph(text)
    _set_rtl(paragraph)
    paragraph.style = document.styles["Normal"]
    for run in paragraph.runs:
        run.font.name = "DejaVu Sans"
        run.font.size = Pt(12)
    return paragraph


def _fixture(scenario: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    if scenario == "explicit_page_break":
        _add_anchor(document, "ANCHOR_PAGE1")
    elif scenario == "multiple_sections":
        _add_anchor(document, "ANCHOR_PAGE1")
    else:
        _add_anchor(document, "ANCHOR_TOP")

    if scenario == "rtl_paragraph":
        _arabic_paragraph(
            document,
            "هاذا التقرير يعرض نتائج المشروع بصورة واضحة ومباشرة.",
        )

    elif scenario == "heading_and_body":
        heading = document.add_heading("هاذا العنوان التنفيذي", level=1)
        _set_rtl(heading)
        _arabic_paragraph(document, "المتن التالي ثابت بعد العنوان.")

    elif scenario == "table_grid":
        table = document.add_table(rows=3, cols=3)
        table.style = "Table Grid"
        table.cell(0, 0).text = "البند"
        table.cell(0, 1).text = "الحالة"
        table.cell(0, 2).text = "القيمة"
        table.cell(1, 0).text = "هاذا البيان"
        table.cell(1, 1).text = "معتمد"
        table.cell(1, 2).text = "100"
        table.cell(2, 0).text = "بيان ثابت"
        table.cell(2, 1).text = "مكتمل"
        table.cell(2, 2).text = "200"

    elif scenario == "merged_table":
        table = document.add_table(rows=3, cols=3)
        table.style = "Table Grid"
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "عنوان مدمج"
        table.cell(0, 2).text = "قيمة"
        table.cell(1, 0).text = "هاذا البيان"
        table.cell(1, 1).text = "ثابت"
        table.cell(1, 2).text = "100"
        table.cell(2, 0).text = "صف ثابت"
        table.cell(2, 1).text = "ثابت"
        table.cell(2, 2).text = "200"

    elif scenario == "numbered_list":
        document.add_paragraph("البند الأول", style="List Number")
        document.add_paragraph("هاذا البند يحتاج تصحيحًا.", style="List Number")
        document.add_paragraph("البند الثالث", style="List Number")

    elif scenario == "bullet_list":
        document.add_paragraph("النقطة الأولى", style="List Bullet")
        document.add_paragraph("هاذا البند يحتاج تصحيحًا.", style="List Bullet")
        document.add_paragraph("النقطة الثالثة", style="List Bullet")

    elif scenario == "image_position":
        paragraph = document.add_paragraph("صورة ثابتة")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(
            BytesIO(_bmp_1x1()),
            width=Inches(1.4),
        )
        _arabic_paragraph(document, "هاذا التقرير أسفل الصورة.")

    elif scenario == "image_caption":
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(
            BytesIO(_bmp_1x1()),
            width=Inches(1.1),
        )
        caption = document.add_paragraph("شكل 1: صورة توضيحية ثابتة")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _arabic_paragraph(document, "هاذا التقرير بعد التسمية.")

    elif scenario == "header_footer":
        section.header.paragraphs[0].text = "VISUAL HEADER"
        section.footer.paragraphs[0].text = "VISUAL FOOTER"
        _arabic_paragraph(document, "هاذا التقرير داخل الصفحة.")

    elif scenario == "landscape_section":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = (
            section.page_height,
            section.page_width,
        )
        _arabic_paragraph(
            document,
            "هاذا التقرير في صفحة أفقية ذات تخطيط ثابت.",
        )

    elif scenario == "explicit_page_break":
        _arabic_paragraph(document, "هاذا التقرير في الصفحة الأولى.")
        paragraph = document.add_paragraph()
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        _add_anchor(document, "ANCHOR_PAGE2")
        _arabic_paragraph(document, "نص ثابت في الصفحة الثانية.")
        return _save(document)

    elif scenario == "multiple_sections":
        _arabic_paragraph(document, "هاذا التقرير في القسم الأول.")
        second = document.add_section(WD_SECTION_START.NEW_PAGE)
        second.top_margin = Inches(1.0)
        second.header.paragraphs[0].text = "SECOND SECTION HEADER"
        _add_anchor(document, "ANCHOR_PAGE2")
        _arabic_paragraph(document, "نص ثابت في القسم الثاني.")
        return _save(document)

    elif scenario == "near_page_boundary":
        for index in range(28):
            _arabic_paragraph(
                document,
                f"فقرة ثابتة رقم {index + 1} لاختبار استقرار توزيع الصفحة.",
            )
        _arabic_paragraph(
            document,
            "هاذا التقرير قرب نهاية الصفحة ويجب ألا يغيّر توزيع الصفحات.",
        )

    elif scenario == "mixed_run_formatting":
        paragraph = document.add_paragraph()
        _set_rtl(paragraph)
        bold = paragraph.add_run("هاذا")
        bold.bold = True
        italic = paragraph.add_run(" التقرير")
        italic.italic = True
        underlined = paragraph.add_run(" معتمد.")
        underlined.underline = True

    elif scenario == "cross_run_fix":
        paragraph = document.add_paragraph()
        _set_rtl(paragraph)
        first = paragraph.add_run("ها")
        first.bold = True
        second = paragraph.add_run("ذا")
        second.italic = True
        paragraph.add_run(" التقرير معتمد.")

    elif scenario == "repeated_fixes":
        _arabic_paragraph(
            document,
            "هاذا التقرير الأول، وهاذا التقرير الثاني، وهاذا الملخص.",
        )

    elif scenario == "bilingual_url":
        paragraph = _arabic_paragraph(
            document,
            "هاذا التقرير يحتوي رابطًا مرجعيًا ثابتًا:",
        )
        paragraph.add_run(" https://example.com/report?id=2026 ")
        _add_hyperlink(
            paragraph,
            "OPENAI-LINK",
            "https://example.com/visual",
        )

    elif scenario == "rich_document":
        section.header.paragraphs[0].text = "RICH HEADER"
        section.footer.paragraphs[0].text = "RICH FOOTER"
        heading = document.add_heading("عنوان تنفيذي ثابت", level=1)
        _set_rtl(heading)
        document.add_paragraph("بند ثابت", style="List Number")
        table = document.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "البند"
        table.cell(0, 1).text = "القيمة"
        table.cell(1, 0).text = "ثابت"
        table.cell(1, 1).text = "100"
        picture = document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(
            BytesIO(_bmp_1x1()),
            width=Inches(0.9),
        )
        _arabic_paragraph(
            document,
            "هاذا التقرير النهائي ضمن مستند غني بالعناصر.",
        )

    elif scenario in {
        "second_round_visual_identity",
        "no_op_visual_identity",
    }:
        text = (
            "هذا التقرير معتمد وخالٍ من التصحيحات الآلية."
            if scenario == "no_op_visual_identity"
            else "هاذا التقرير معتمد."
        )
        _arabic_paragraph(document, text)

    else:
        raise ValueError(f"Unknown scenario: {scenario}")

    _add_anchor(document, "ANCHOR_BOTTOM")
    return _save(document)


def _render_docx(data: bytes, workdir: Path, stem: str) -> Path:
    soffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not soffice:
        raise RuntimeError("LibreOffice/soffice is required.")

    source_dir = workdir / f"{stem}-src"
    output_dir = workdir / f"{stem}-pdf"
    profile_dir = workdir / f"{stem}-profile"
    source_dir.mkdir()
    output_dir.mkdir()
    profile_dir.mkdir()

    docx_path = source_dir / f"{stem}.docx"
    docx_path.write_bytes(data)

    profile_uri = profile_dir.resolve().as_uri()
    command = [
        soffice,
        "--headless",
        f"-env:UserInstallation={profile_uri}",
        "--convert-to",
        "pdf:writer_pdf_Export",
        "--outdir",
        str(output_dir),
        str(docx_path),
    ]
    result = subprocess.run(
        command,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        timeout=90,
        check=False,
    )
    pdf_path = output_dir / f"{stem}.pdf"
    if result.returncode != 0 or not pdf_path.exists():
        raise RuntimeError(
            "LibreOffice render failed: "
            + result.stdout
            + "\n"
            + result.stderr
        )
    return pdf_path


def _analyze_pdf(path: Path):
    document = fitz.open(path)
    pages = []

    for page in document:
        pix = page.get_pixmap(
            matrix=fitz.Matrix(1.5, 1.5),
            alpha=False,
        )
        image = Image.frombytes(
            "RGB",
            [pix.width, pix.height],
            pix.samples,
        )

        image_rects = []
        for image_info in page.get_image_info(xrefs=True):
            bbox = image_info.get("bbox")
            if bbox:
                image_rects.append(tuple(float(x) for x in bbox))

        drawings = []
        for drawing in page.get_drawings():
            rect = drawing.get("rect")
            if rect:
                drawings.append(
                    (
                        round(float(rect.x0), 2),
                        round(float(rect.y0), 2),
                        round(float(rect.x1), 2),
                        round(float(rect.y1), 2),
                    )
                )

        pages.append(
            {
                "width": float(page.rect.width),
                "height": float(page.rect.height),
                "image": image,
                "image_rects": sorted(image_rects),
                "drawings": sorted(drawings),
            }
        )

    return document, pages


def _rect_distance(left, right) -> float:
    return max(
        abs(float(a) - float(b))
        for a, b in zip(left, right)
    )


def _find_anchor(pdf: fitz.Document, text: str):
    hits = []
    for page_index, page in enumerate(pdf):
        for rect in page.search_for(text):
            hits.append(
                (
                    page_index,
                    (
                        float(rect.x0),
                        float(rect.y0),
                        float(rect.x1),
                        float(rect.y1),
                    ),
                )
            )
    return hits


def _changed_pixel_ratio(left: Image.Image, right: Image.Image) -> float:
    if left.size != right.size:
        return 1.0
    diff = ImageChops.difference(left, right).convert("L")
    histogram = diff.histogram()
    changed = sum(
        count
        for value, count in enumerate(histogram)
        if value >= 12
    )
    total = left.width * left.height
    return changed / total if total else 0.0


def _apply_nadid(data: bytes):
    nodes = parse_docx(data)
    protected = extract_protected_spans(nodes)
    suggestions = fast_review(nodes)
    plan = build_safe_auto_apply_plan(
        nodes,
        suggestions,
        protected,
    )
    output, report = apply_patches_to_docx(
        data,
        plan.patches,
    )
    return output, report, plan


def _compare_visual(
    before_pdf: Path,
    after_pdf: Path,
    case: dict,
    defaults: dict,
) -> tuple[list[str], dict]:
    failures = []
    metrics = {}

    before_doc, before_pages = _analyze_pdf(before_pdf)
    after_doc, after_pages = _analyze_pdf(after_pdf)

    if len(before_pages) != len(after_pages):
        failures.append(
            f"page_count:{len(before_pages)}!={len(after_pages)}"
        )

    expected_pages = case.get("expected_pages")
    if (
        expected_pages is not None
        and len(after_pages) != expected_pages
    ):
        failures.append(
            f"expected_pages:{len(after_pages)}!={expected_pages}"
        )

    page_pairs = min(len(before_pages), len(after_pages))
    ratios = []
    for index in range(page_pairs):
        before = before_pages[index]
        after = after_pages[index]

        if (
            abs(before["width"] - after["width"]) > 0.5
            or abs(before["height"] - after["height"]) > 0.5
        ):
            failures.append(f"page_geometry_changed:{index + 1}")

        ratio = _changed_pixel_ratio(
            before["image"],
            after["image"],
        )
        ratios.append(ratio)

    max_ratio = max(ratios) if ratios else 0.0
    metrics["max_changed_pixel_ratio"] = round(max_ratio, 6)
    allowed_ratio = float(
        case.get(
            "max_changed_pixel_ratio",
            defaults["max_changed_pixel_ratio"],
        )
    )
    if max_ratio > allowed_ratio:
        failures.append(
            f"visual_diff_ratio:{max_ratio:.4f}>{allowed_ratio:.4f}"
        )

    max_anchor_drift = float(defaults["max_anchor_drift_pt"])
    anchor_metrics = {}
    for anchor in case.get("anchors", []):
        before_hits = _find_anchor(before_doc, anchor)
        after_hits = _find_anchor(after_doc, anchor)
        if len(before_hits) != len(after_hits) or not before_hits:
            failures.append(f"anchor_missing_or_changed:{anchor}")
            continue

        drifts = []
        for left, right in zip(before_hits, after_hits):
            if left[0] != right[0]:
                failures.append(f"anchor_page_changed:{anchor}")
                continue
            drift = _rect_distance(left[1], right[1])
            drifts.append(drift)
            if drift > max_anchor_drift:
                failures.append(
                    f"anchor_drift:{anchor}:{drift:.2f}"
                )
        if drifts:
            anchor_metrics[anchor] = round(max(drifts), 3)

    metrics["anchor_drift_pt"] = anchor_metrics

    if case.get("check_images"):
        before_rects = [
            rect
            for page in before_pages
            for rect in page["image_rects"]
        ]
        after_rects = [
            rect
            for page in after_pages
            for rect in page["image_rects"]
        ]
        if len(before_rects) != len(after_rects):
            failures.append(
                f"image_count:{len(before_rects)}!={len(after_rects)}"
            )
        else:
            allowed = float(defaults["max_graphic_drift_pt"])
            for index, (left, right) in enumerate(
                zip(before_rects, after_rects)
            ):
                drift = _rect_distance(left, right)
                if drift > allowed:
                    failures.append(
                        f"image_drift:{index}:{drift:.2f}"
                    )

    if case.get("check_drawings"):
        before_drawings = [
            rect
            for page in before_pages
            for rect in page["drawings"]
        ]
        after_drawings = [
            rect
            for page in after_pages
            for rect in page["drawings"]
        ]
        if len(before_drawings) != len(after_drawings):
            failures.append(
                "drawing_count:"
                f"{len(before_drawings)}!={len(after_drawings)}"
            )
        else:
            allowed = float(defaults["max_graphic_drift_pt"])
            for index, (left, right) in enumerate(
                zip(before_drawings, after_drawings)
            ):
                drift = _rect_distance(left, right)
                if drift > allowed:
                    failures.append(
                        f"drawing_drift:{index}:{drift:.2f}"
                    )

    before_doc.close()
    after_doc.close()
    return failures, metrics


def evaluate(case: dict, defaults: dict) -> dict:
    source = _fixture(case["scenario"])

    if case.get("no_op"):
        output = source
        applied = 0
    else:
        output, report, plan = _apply_nadid(source)
        if not report.fidelity_ok:
            return {
                "id": case["id"],
                "scenario": case["scenario"],
                "passed": False,
                "failures": [
                    "document_fidelity_failed:"
                    + ",".join(report.fidelity_errors or [])
                ],
            }
        applied = len(report.applied)
        if not plan.patches:
            return {
                "id": case["id"],
                "scenario": case["scenario"],
                "passed": False,
                "failures": ["no_auto_fix_generated"],
            }

    with tempfile.TemporaryDirectory(prefix="nadid-visual-") as temp:
        workdir = Path(temp)
        before_pdf = _render_docx(
            source,
            workdir,
            "before",
        )
        after_pdf = _render_docx(
            output,
            workdir,
            "after",
        )
        failures, metrics = _compare_visual(
            before_pdf,
            after_pdf,
            case,
            defaults,
        )

        if case.get("second_round_identical"):
            second_output, second_report, second_plan = _apply_nadid(
                output
            )
            if second_plan.patches:
                failures.append("second_round_has_auto_fix")
            if second_output != output:
                failures.append("second_round_docx_changed")
            if not second_report.fidelity_ok:
                failures.append("second_round_fidelity_failed")

            second_pdf = _render_docx(
                second_output,
                workdir,
                "second",
            )
            second_failures, second_metrics = _compare_visual(
                after_pdf,
                second_pdf,
                {
                    **case,
                    "max_changed_pixel_ratio": 0.0,
                },
                defaults,
            )
            failures.extend(
                f"second:{item}"
                for item in second_failures
            )
            metrics["second_round"] = second_metrics

    return {
        "id": case["id"],
        "scenario": case["scenario"],
        "passed": not failures,
        "failures": failures,
        "applied": applied,
        "metrics": metrics,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--suite",
        default="benchmark/visual/v1.json",
    )
    parser.add_argument("--report")
    parser.add_argument("--enforce", action="store_true")
    args = parser.parse_args()

    if not (shutil.which("libreoffice") or shutil.which("soffice")):
        raise SystemExit(
            "LibreOffice/soffice is required for visual regression."
        )

    payload = json.loads(
        Path(args.suite).read_text(encoding="utf-8")
    )
    defaults = payload["defaults"]
    results = [
        evaluate(case, defaults)
        for case in payload["cases"]
    ]

    passed = sum(item["passed"] for item in results)
    total = len(results)
    score = passed / total if total else 0.0

    report = {
        "schema_version": 1,
        "suite": args.suite,
        "renderer": "libreoffice_writer_pdf_export",
        "passed": passed,
        "total": total,
        "score": round(score, 4),
        "failures": [
            item for item in results if not item["passed"]
        ],
        "results": results,
    }

    print(
        "Visual Layout & Rendering Regression V1: "
        f"{passed}/{total} = {score:.1%}"
    )

    if report["failures"]:
        print("\nFailures:")
        for item in report["failures"]:
            print(
                f"- {item['id']}: "
                + ", ".join(item["failures"])
            )

    if args.report:
        Path(args.report).write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    if args.enforce and report["failures"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
