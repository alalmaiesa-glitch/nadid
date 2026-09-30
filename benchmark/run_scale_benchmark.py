from __future__ import annotations

import argparse
import gc
import json
import platform
import time
import tracemalloc
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document

from app.pipeline.chunker import build_chunks
from app.pipeline.memory import build_document_memory
from app.pipeline.parser import parse_docx
from app.pipeline.protection import extract_protected_spans
from app.pipeline.reviewer import fast_review


ROOT = Path(__file__).resolve().parent
DEFAULT_CONFIG = ROOT / "scale" / "v1.json"


def _build_fixture(paragraph_count: int) -> bytes:
    document = Document()
    document.add_heading("اختبار نَضِيد للمستندات العربية الطويلة", level=1)

    for index in range(paragraph_count):
        if index and index % 50 == 0:
            document.add_heading(
                f"القسم التشغيلي {index // 50 + 1}",
                level=2,
            )

        budget = 1000 + (index % 17)
        document.add_paragraph(
            "يعرض هاذا القسم حالة المشروع بصورة موجزة، "
            f"والبند رقم {index + 1} ضمن خطة عام 2026، "
            f"وتبلغ القيمة المرجعية {budget} ريال، "
            "ويجب أن يلتزم الفريق بالمدة والنطاق المعتمدين."
        )

        if index and index % 250 == 0:
            table = document.add_table(rows=3, cols=3)
            table.style = "Table Grid"
            table.cell(0, 0).text = "البند"
            table.cell(0, 1).text = "الحالة"
            table.cell(0, 2).text = "القيمة"
            table.cell(1, 0).text = f"مرحلة {index // 250}"
            table.cell(1, 1).text = "معتمد"
            table.cell(1, 2).text = str(budget)
            table.cell(2, 0).text = "المراجعة"
            table.cell(2, 1).text = "مستمرة"
            table.cell(2, 2).text = "2026"

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _measure(callable_):
    started = time.perf_counter()
    value = callable_()
    return value, time.perf_counter() - started


def _run_case(case: dict[str, Any]) -> dict[str, Any]:
    paragraph_count = int(case["paragraphs"])

    fixture_started = time.perf_counter()
    payload = _build_fixture(paragraph_count)
    fixture_seconds = time.perf_counter() - fixture_started

    gc.collect()
    tracemalloc.start()

    nodes, parse_seconds = _measure(lambda: parse_docx(payload))
    chunks, chunk_seconds = _measure(lambda: build_chunks(nodes))
    protected, protection_seconds = _measure(
        lambda: extract_protected_spans(nodes)
    )
    suggestions, review_seconds = _measure(lambda: fast_review(nodes))
    memory, memory_seconds = _measure(
        lambda: build_document_memory(nodes, chunks, protected)
    )

    _, peak_bytes = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    # Re-run deterministic outputs outside the measured window. This is
    # intentionally excluded from the performance budget.
    second_nodes = parse_docx(payload)
    second_suggestions = fast_review(second_nodes)

    node_ids = [item.id for item in nodes]
    second_node_ids = [item.id for item in second_nodes]
    suggestion_ids = [item.id for item in suggestions]
    second_suggestion_ids = [item.id for item in second_suggestions]

    word_count = sum(len(node.text.split()) for node in nodes)
    stage_seconds = {
        "parse": parse_seconds,
        "chunk": chunk_seconds,
        "protection": protection_seconds,
        "review": review_seconds,
        "memory": memory_seconds,
    }
    total_seconds = sum(stage_seconds.values())
    peak_mib = peak_bytes / (1024 * 1024)

    failures: list[str] = []

    if node_ids != second_node_ids:
        failures.append("node_ids_not_deterministic")
    if suggestion_ids != second_suggestion_ids:
        failures.append("suggestion_ids_not_deterministic")
    if memory.chunk_count != len(chunks):
        failures.append("memory_chunk_count_mismatch")
    if memory.protected_count != len(protected):
        failures.append("memory_protected_count_mismatch")
    if not nodes or not chunks:
        failures.append("empty_pipeline_output")

    minimum_ratio = float(case.get("minimum_suggestions_per_paragraph", 0.95))
    if len(suggestions) < int(paragraph_count * minimum_ratio):
        failures.append("unexpected_suggestion_drop")

    max_seconds = float(case["max_total_seconds"])
    if total_seconds > max_seconds:
        failures.append(
            f"total_seconds_budget_exceeded:{total_seconds:.3f}>{max_seconds:.3f}"
        )

    max_peak_mib = float(case["max_peak_mib"])
    if peak_mib > max_peak_mib:
        failures.append(
            f"python_peak_mib_budget_exceeded:{peak_mib:.1f}>{max_peak_mib:.1f}"
        )

    return {
        "id": case["id"],
        "paragraphs": paragraph_count,
        "docx_bytes": len(payload),
        "word_count": word_count,
        "node_count": len(nodes),
        "chunk_count": len(chunks),
        "protected_count": len(protected),
        "suggestion_count": len(suggestions),
        "fact_count": len(memory.facts),
        "knowledge_item_count": len(memory.knowledge_items),
        "fixture_seconds": round(fixture_seconds, 6),
        "stage_seconds": {
            key: round(value, 6)
            for key, value in stage_seconds.items()
        },
        "total_seconds": round(total_seconds, 6),
        "python_peak_mib": round(peak_mib, 3),
        "seconds_per_1000_nodes": round(
            total_seconds / max(len(nodes), 1) * 1000,
            6,
        ),
        "deterministic_nodes": node_ids == second_node_ids,
        "deterministic_suggestions": (
            suggestion_ids == second_suggestion_ids
        ),
        "failures": failures,
        "passed": not failures,
    }


def _growth_checks(
    results: list[dict[str, Any]],
    config: dict[str, Any],
) -> list[dict[str, Any]]:
    growth = config["growth"]
    max_time_growth = float(growth["max_normalized_time_growth"])
    max_peak_growth = float(growth["max_normalized_peak_growth"])

    checks: list[dict[str, Any]] = []
    for previous, current in zip(results, results[1:]):
        node_ratio = current["node_count"] / max(previous["node_count"], 1)
        time_ratio = current["total_seconds"] / max(
            previous["total_seconds"],
            0.001,
        )
        peak_ratio = current["python_peak_mib"] / max(
            previous["python_peak_mib"],
            0.001,
        )

        normalized_time_growth = time_ratio / node_ratio
        normalized_peak_growth = peak_ratio / node_ratio

        failures: list[str] = []
        if normalized_time_growth > max_time_growth:
            failures.append(
                "normalized_time_growth_exceeded:"
                f"{normalized_time_growth:.3f}>{max_time_growth:.3f}"
            )
        if normalized_peak_growth > max_peak_growth:
            failures.append(
                "normalized_peak_growth_exceeded:"
                f"{normalized_peak_growth:.3f}>{max_peak_growth:.3f}"
            )

        checks.append(
            {
                "from": previous["id"],
                "to": current["id"],
                "node_ratio": round(node_ratio, 4),
                "time_ratio": round(time_ratio, 4),
                "peak_ratio": round(peak_ratio, 4),
                "normalized_time_growth": round(
                    normalized_time_growth,
                    4,
                ),
                "normalized_peak_growth": round(
                    normalized_peak_growth,
                    4,
                ),
                "failures": failures,
                "passed": not failures,
            }
        )
    return checks


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--enforce", action="store_true")
    args = parser.parse_args()

    config = json.loads(args.config.read_text(encoding="utf-8"))
    results = [_run_case(case) for case in config["cases"]]
    growth_checks = _growth_checks(results, config)

    failed_cases = [
        result["id"]
        for result in results
        if not result["passed"]
    ]
    failed_growth = [
        f"{check['from']}->{check['to']}"
        for check in growth_checks
        if not check["passed"]
    ]

    report = {
        "schema_version": 1,
        "benchmark": "large_document_scale_performance_v1",
        "environment": {
            "python": platform.python_version(),
            "platform": platform.platform(),
        },
        "measurement_policy": {
            "timing": "perf_counter per pipeline stage",
            "memory": (
                "tracemalloc Python allocation peak; "
                "not whole-process RSS"
            ),
            "fixture_generation_in_budget": False,
            "determinism_rerun_in_budget": False,
        },
        "cases": results,
        "growth_checks": growth_checks,
        "summary": {
            "passed_cases": len(results) - len(failed_cases),
            "total_cases": len(results),
            "failed_cases": failed_cases,
            "failed_growth_checks": failed_growth,
            "passed": not failed_cases and not failed_growth,
        },
    }

    rendered = json.dumps(
        report,
        ensure_ascii=False,
        indent=2,
        sort_keys=True,
    )
    print(rendered)

    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(rendered + "\n", encoding="utf-8")

    if args.enforce and not report["summary"]["passed"]:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
