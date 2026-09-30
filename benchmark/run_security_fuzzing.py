from __future__ import annotations

import argparse
import json
import random
import time
import warnings
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from docx import Document
from fastapi.testclient import TestClient

import app.docx_security as security
from app.docx_security import UnsafeDocxError, validate_docx_payload
from app.main import app


ROOT = Path(__file__).resolve().parent
DEFAULT_CONFIG = ROOT / "security" / "v1.json"
CLIENT = TestClient(app, raise_server_exceptions=False)

REL_NS = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
HYPERLINK_REL = (
    "http://schemas.openxmlformats.org/officeDocument/"
    "2006/relationships/hyperlink"
)
IMAGE_REL = (
    "http://schemas.openxmlformats.org/officeDocument/"
    "2006/relationships/image"
)
TEMPLATE_REL = (
    "http://schemas.openxmlformats.org/officeDocument/"
    "2006/relationships/attachedTemplate"
)
OLE_REL = (
    "http://schemas.openxmlformats.org/officeDocument/"
    "2006/relationships/oleObject"
)


def _valid_docx() -> bytes:
    document = Document()
    document.add_heading("اختبار أمان OOXML", level=1)
    document.add_paragraph(
        "هاذا المستند المرجعي آمن ويجب أن يمر عبر نَضِيد."
    )
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _rewrite(
    payload: bytes,
    replacements: dict[str, bytes] | None = None,
    additions: list[tuple[str, bytes]] | None = None,
) -> bytes:
    replacements = replacements or {}
    additions = additions or []
    output = BytesIO()

    with ZipFile(BytesIO(payload), "r") as source:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            with ZipFile(
                output,
                "w",
                compression=ZIP_DEFLATED,
            ) as target:
                for info in source.infolist():
                    target.writestr(
                        info.filename,
                        replacements.get(
                            info.filename,
                            source.read(info.filename),
                        ),
                    )
                for name, value in additions:
                    target.writestr(name, value)

    return output.getvalue()


def _inject_relationship(
    payload: bytes,
    rel_type: str,
    target: str,
    target_mode: str | None = "External",
) -> bytes:
    rel_name = "word/_rels/document.xml.rels"
    with ZipFile(BytesIO(payload), "r") as archive:
        rels = archive.read(rel_name)

    attributes = (
        'Id="rIdSecurityFuzz" '
        f'Type="{rel_type}" Target="{target}"'
    )
    if target_mode is not None:
        attributes += f' TargetMode="{target_mode}"'

    relation = f"<Relationship {attributes}/>".encode("utf-8")
    marker = b"</Relationships>"

    return _rewrite(
        payload,
        replacements={
            rel_name: rels.replace(
                marker,
                relation + marker,
            )
        },
    )


def _replace_content_type(payload: bytes) -> bytes:
    name = "[Content_Types].xml"
    with ZipFile(BytesIO(payload), "r") as archive:
        content = archive.read(name)
    content = content.replace(
        b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
        b"application/vnd.ms-word.document.macroEnabled.main+xml",
    )
    return _rewrite(payload, replacements={name: content})


def _fixture(kind: str, valid: bytes) -> bytes:
    if kind == "valid":
        return valid
    if kind == "https_hyperlink":
        return _inject_relationship(
            valid,
            HYPERLINK_REL,
            "https://example.com/reference",
        )
    if kind == "mailto_hyperlink":
        return _inject_relationship(
            valid,
            HYPERLINK_REL,
            "mailto:editor@example.com",
        )
    if kind == "file_hyperlink":
        return _inject_relationship(
            valid,
            HYPERLINK_REL,
            "file:///etc/passwd",
        )
    if kind == "ftp_hyperlink":
        return _inject_relationship(
            valid,
            HYPERLINK_REL,
            "ftp://example.com/reference",
        )
    if kind == "javascript_hyperlink":
        return _inject_relationship(
            valid,
            HYPERLINK_REL,
            "javascript:alert(1)",
        )
    if kind == "external_template":
        return _inject_relationship(
            valid,
            TEMPLATE_REL,
            "https://example.com/template.dotm",
        )
    if kind == "external_image":
        return _inject_relationship(
            valid,
            IMAGE_REL,
            "https://example.com/image.png",
        )
    if kind == "internal_ole_relationship":
        return _inject_relationship(
            valid,
            OLE_REL,
            "embeddings/oleObject1.bin",
            target_mode=None,
        )
    if kind == "vba_project":
        return _rewrite(
            valid,
            additions=[("word/vbaProject.bin", b"macro")],
        )
    if kind == "vba_signature":
        return _rewrite(
            valid,
            additions=[("word/vbaProjectSignature.bin", b"sig")],
        )
    if kind == "activex":
        return _rewrite(
            valid,
            additions=[("word/activeX/activeX1.bin", b"activex")],
        )
    if kind == "ole_embedding":
        return _rewrite(
            valid,
            additions=[("word/embeddings/oleObject1.bin", b"ole")],
        )
    if kind == "custom_ui":
        return _rewrite(
            valid,
            additions=[("customUI/customUI.xml", b"<customUI/>")],
        )
    if kind == "macro_content_type":
        return _replace_content_type(valid)
    if kind == "xml_doctype":
        xml = b"""<?xml version="1.0"?>
<!DOCTYPE w:document [<!ENTITY x "expanded">]>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body><w:p><w:r><w:t>&x;</w:t></w:r></w:p></w:body>
</w:document>"""
        return _rewrite(
            valid,
            replacements={"word/document.xml": xml},
        )
    if kind == "xml_entity":
        xml = b"""<?xml version="1.0"?>
<!DOCTYPE root [<!ENTITY a "A">]>
<root>&a;</root>"""
        return _rewrite(
            valid,
            replacements={"word/document.xml": xml},
        )
    if kind == "deep_xml":
        xml = (
            b"<root>"
            + b"<n>" * 40
            + b"x"
            + b"</n>" * 40
            + b"</root>"
        )
        return _rewrite(
            valid,
            replacements={"word/document.xml": xml},
        )
    if kind == "many_xml_elements":
        xml = b"<root>" + b"<n/>" * 100 + b"</root>"
        return _rewrite(
            valid,
            replacements={"word/document.xml": xml},
        )
    if kind == "many_xml_attributes":
        attrs = b" ".join(
            f'a{i}="{i}"'.encode("ascii")
            for i in range(30)
        )
        xml = b"<root " + attrs + b"/>"
        return _rewrite(
            valid,
            replacements={"word/document.xml": xml},
        )
    if kind == "large_xml_entry":
        xml = b"<root>" + b"x" * 4096 + b"</root>"
        return _rewrite(
            valid,
            replacements={"word/document.xml": xml},
        )
    if kind == "large_total_xml":
        additions = [
            (
                f"word/security-{index}.xml",
                b"<root>" + b"x" * 700 + b"</root>",
            )
            for index in range(5)
        ]
        return _rewrite(valid, additions=additions)
    if kind == "large_binary_entry":
        return _rewrite(
            valid,
            additions=[("word/media/security.bin", b"x" * 4096)],
        )
    if kind == "compressible_entry":
        return _rewrite(
            valid,
            additions=[("word/media/compressible.bin", b"A" * 20_000)],
        )
    if kind == "many_entries":
        return _rewrite(
            valid,
            additions=[
                (f"security/e-{index}.bin", b"x")
                for index in range(100)
            ],
        )
    if kind == "duplicate_entry":
        return _rewrite(
            valid,
            additions=[("word/document.xml", b"<duplicate/>")],
        )
    if kind == "path_traversal":
        return _rewrite(
            valid,
            additions=[("../security.txt", b"x")],
        )
    if kind == "corrupt_relationship_xml":
        name = "word/_rels/document.xml.rels"
        return _rewrite(
            valid,
            replacements={name: b"<Relationships"},
        )
    if kind == "corrupt_content_types_xml":
        return _rewrite(
            valid,
            replacements={"[Content_Types].xml": b"<Types"},
        )
    raise KeyError(kind)


@contextmanager
def _temporary_guard(name: str | None, value: Any):
    if not name:
        yield
        return
    original = getattr(security, name)
    setattr(security, name, value)
    try:
        yield
    finally:
        setattr(security, name, original)


def _post(endpoint: str, payload: bytes):
    path = {
        "analyze": "/v1/analyze/docx",
        "deep": "/v1/analyze/docx/deep",
        "apply": "/v1/apply/docx",
    }[endpoint]
    data = {"patches": "[]"} if endpoint == "apply" else {}
    return CLIENT.post(
        path,
        files={
            "file": (
                "security.docx",
                payload,
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            )
        },
        data=data,
    )


def _detail(response) -> str:
    if response.status_code == 200:
        return ""
    try:
        body = response.json()
    except Exception:
        return response.text[:500]
    if isinstance(body, dict):
        body = body.get("detail", body)
    return json.dumps(body, ensure_ascii=False, sort_keys=True)[:500]


def _run_case(case: dict[str, Any], valid: bytes) -> dict[str, Any]:
    payload = _fixture(case["fixture"], valid)
    failures: list[str] = []
    started = time.perf_counter()

    if case["mode"] == "endpoint":
        response = _post(case["endpoint"], payload)
        status = response.status_code
        detail = _detail(response)

        if status != int(case["expected_status"]):
            failures.append(
                f"status:{status}!={case['expected_status']}"
            )
        expected = case.get("detail_contains")
        if expected and expected not in detail:
            failures.append("missing_expected_detail")
        if status >= 500:
            failures.append("server_error_escape")

        observed = {
            "status": status,
            "detail": detail,
        }
    else:
        error = None
        with _temporary_guard(
            case.get("guard"),
            case.get("guard_value"),
        ):
            try:
                validate_docx_payload(payload)
            except UnsafeDocxError as exc:
                error = str(exc)

        if error != case["expected_error"]:
            failures.append(
                f"error:{error}!={case['expected_error']}"
            )
        observed = {"error": error}

    elapsed = time.perf_counter() - started
    return {
        "id": case["id"],
        "fixture": case["fixture"],
        "mode": case["mode"],
        "payload_bytes": len(payload),
        "elapsed_seconds": round(elapsed, 6),
        "observed": observed,
        "failures": failures,
        "passed": not failures,
    }


def _mutate_payloads(
    payload: bytes,
    seed: int,
    variants: int,
):
    rng = random.Random(seed)

    truncations = [
        0.08,
        0.18,
        0.33,
        0.50,
        0.72,
        0.90,
    ]
    generated: list[tuple[str, bytes]] = []

    for ratio in truncations:
        end = max(1, int(len(payload) * ratio))
        generated.append((f"truncate-{ratio:.2f}", payload[:end]))

    while len(generated) < variants:
        mutated = bytearray(payload)
        flips = 1 + (len(generated) % 3)

        for _ in range(flips):
            low = min(32, len(mutated) - 1)
            high = max(low, len(mutated) - 1)
            index = rng.randint(low, high)
            mutated[index] ^= 1 << rng.randint(0, 7)

        generated.append(
            (f"bitflip-{len(generated) + 1}", bytes(mutated))
        )

    return generated[:variants]


def _run_mutation_fuzz(
    payload: bytes,
    config: dict[str, Any],
) -> dict[str, Any]:
    allowed = set(config["allowed_http_statuses"])
    max_seconds = float(config["max_case_seconds"])
    rows = []
    failures = []

    for name, mutant in _mutate_payloads(
        payload,
        int(config["seed"]),
        int(config["variants"]),
    ):
        started = time.perf_counter()
        response = _post("analyze", mutant)
        elapsed = time.perf_counter() - started
        status = response.status_code
        row_failures = []

        if status not in allowed:
            row_failures.append(f"unexpected_status:{status}")
        if status >= 500:
            row_failures.append("server_error_escape")
        if elapsed > max_seconds:
            row_failures.append(
                f"case_time:{elapsed:.3f}>{max_seconds:.3f}"
            )

        rows.append(
            {
                "name": name,
                "payload_bytes": len(mutant),
                "status": status,
                "elapsed_seconds": round(elapsed, 6),
                "failures": row_failures,
                "passed": not row_failures,
            }
        )
        if row_failures:
            failures.append(name)

    return {
        "seed": config["seed"],
        "variants": len(rows),
        "allowed_http_statuses": sorted(allowed),
        "failed_variants": failures,
        "server_errors": sum(
            1 for row in rows if row["status"] >= 500
        ),
        "accepted_variants": sum(
            1 for row in rows if row["status"] == 200
        ),
        "rejected_variants": sum(
            1 for row in rows if row["status"] != 200
        ),
        "passed": not failures,
        "cases": rows,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--enforce", action="store_true")
    args = parser.parse_args()

    config = json.loads(args.config.read_text(encoding="utf-8"))
    valid = _valid_docx()
    cases = [_run_case(case, valid) for case in config["cases"]]
    mutation = _run_mutation_fuzz(
        valid,
        config["mutation_fuzz"],
    )

    failed_cases = [
        case["id"] for case in cases if not case["passed"]
    ]
    passed = not failed_cases and mutation["passed"]

    report = {
        "schema_version": 1,
        "benchmark": "adversarial_ooxml_resource_security_fuzzing_v1",
        "policy": {
            "active_content": "reject",
            "xml_dtd_entities": "reject",
            "external_relationships": (
                "allow only http/https/mailto hyperlinks"
            ),
            "resource_fuzzing": (
                "bounded fixtures with reduced thresholds"
            ),
        },
        "cases": cases,
        "mutation_fuzz": mutation,
        "summary": {
            "passed": passed,
            "passed_cases": len(cases) - len(failed_cases),
            "total_cases": len(cases),
            "failed_cases": failed_cases,
            "mutation_variants": mutation["variants"],
            "mutation_failed_variants": mutation["failed_variants"],
            "server_errors": (
                sum(
                    1
                    for case in cases
                    if case["observed"].get("status", 0) in range(500, 600)
                )
                + mutation["server_errors"]
            ),
        },
    }

    rendered = json.dumps(
        report,
        ensure_ascii=False,
        indent=2,
        sort_keys=True,
    )
    print(rendered)

    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(rendered + "\n", encoding="utf-8")

    if args.enforce and not passed:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
