from __future__ import annotations

import argparse
import json
from io import BytesIO
from pathlib import Path

from docx import Document

from app.pipeline.chunker import build_chunks
from app.pipeline.memory import build_document_memory
from app.pipeline.parser import parse_docx
from app.pipeline.protection import extract_protected_spans
from app.pipeline.review_loop import run_docx_review_loop
from app.pipeline.semantic_review import semantic_review


def _save(document: Document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _plain_doc(paragraphs: list[str]) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    return _save(document)


def build_docx(scenario: str) -> bytes:
    if scenario == "clean_document":
        return _plain_doc([
            "تم اعتماد الخطة وفق نتائج الدراسة، وسيبدأ التنفيذ لاحقًا.",
            "هذا التقرير مكتمل ومراجع.",
        ])

    if scenario == "single_orthography":
        return _plain_doc(["هاذا التقرير معتمد."])

    if scenario == "single_punctuation":
        return _plain_doc(["تم اعتماد الخطة ،ثم بدأ التنفيذ."])

    if scenario == "dense_errors":
        return _plain_doc([
            "هاذا التقرير ،ثم بدأ التنفيذ.. لاكن الخطة معتمدة."
        ])

    if scenario == "multi_paragraph":
        return _plain_doc([
            "هاذا التقرير الأول.",
            "لاكن الخطة الثانية معتمدة.",
            "هذة الفقرة الثالثة مكتملة.",
        ])

    if scenario == "heading_fix":
        document = Document()
        document.add_heading("هاذا العنوان", level=1)
        document.add_paragraph("المتن سليم.")
        return _save(document)

    if scenario == "table_cell_fix":
        document = Document()
        document.add_paragraph("مقدمة سليمة.")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "هاذا البيان"
        return _save(document)

    if scenario == "critical_clause_punctuation":
        return _plain_doc([
            "لا يتحمل الطرف الأول تكاليف النقل ، وفق العقد."
        ])

    if scenario == "protected_values_punctuation":
        return _plain_doc([
            "بلغت القيمة 100 ريال ، في 15/10/2026."
        ])

    if scenario == "style_only":
        return _plain_doc([
            "تمت المعالجة في الوقت الراهن وفق الإجراء المعتمد."
        ])

    if scenario == "mixed_style_language":
        return _plain_doc([
            "هاذا التقرير في الوقت الراهن قيد المراجعة."
        ])

    if scenario == "decision_conflict_with_fix":
        return _plain_doc([
            "تم اعتماد الخطة التشغيلية للمشروع.",
            "تم إلغاء الخطة التشغيلية للمشروع.",
            "هاذا ملخص القرارين.",
        ])

    if scenario == "fact_conflict_with_fix":
        return _plain_doc([
            "بلغت الطاقة التشغيلية السنوية للمصنع 2000000 وحدة.",
            "إجمالي القدرة الإنتاجية للمصنع سنويًا تساوي 1800000 وحدة.",
            "هاذا ملخص الأرقام.",
        ])

    if scenario in {
        "second_run_idempotent",
        "one_round_converges",
        "clean_second_run_bytes",
    }:
        return _plain_doc([
            "هاذا التقرير ،ثم بدأ التنفيذ."
        ])

    if scenario == "repeated_occurrences":
        return _plain_doc([
            "هاذا التقرير ،ثم لاكن التنفيذ مستمر ،ثم هاذا الملخص."
        ])

    if scenario == "formatted_run":
        document = Document()
        paragraph = document.add_paragraph()
        first = paragraph.add_run("هاذا")
        first.bold = True
        paragraph.add_run(" التقرير معتمد.")
        return _save(document)

    if scenario == "style_and_punctuation_same_node":
        return _plain_doc([
            "تمت المعالجة في الوقت الراهن ،ثم استمر التنفيذ."
        ])

    if scenario == "semantic_only":
        return _plain_doc([
            "المشروع معتمد وفق المحضر النهائي.",
            "المشروع غير معتمد وفق النسخة الأخيرة.",
        ])

    raise ValueError(f"Unknown scenario: {scenario}")


def _all_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _semantic_issue_types(data: bytes) -> set[str]:
    nodes = parse_docx(data)
    chunks = build_chunks(nodes)
    protected = extract_protected_spans(nodes)
    memory = build_document_memory(nodes, chunks, protected)
    return {
        issue.issue_type
        for issue in semantic_review(nodes, memory)
    }


def _fact_conflict_count(data: bytes) -> int:
    nodes = parse_docx(data)
    chunks = build_chunks(nodes)
    protected = extract_protected_spans(nodes)
    memory = build_document_memory(nodes, chunks, protected)
    return len(memory.conflicts)


def _applied_count(result) -> int:
    return sum(item.applied_count for item in result.rounds)


def evaluate(case: dict) -> dict:
    scenario = case["scenario"]
    expected = case.get("expect", {})
    source = build_docx(scenario)
    result = run_docx_review_loop(
        source,
        max_rounds=int(case.get("max_rounds", 4)),
    )
    failures: list[str] = []

    scalar = {
        "stable": result.stable,
        "safety": result.safety_pass,
        "meaning": result.meaning_preserved,
    }
    for key, actual in scalar.items():
        if key in expected and actual != expected[key]:
            failures.append(f"{key}:{actual}!={expected[key]}")

    if "initial_auto" in expected:
        if result.initial_auto_fix_count != expected["initial_auto"]:
            failures.append(
                f"initial_auto:{result.initial_auto_fix_count}"
            )

    if "min_initial_auto" in expected:
        if result.initial_auto_fix_count < expected["min_initial_auto"]:
            failures.append(
                f"initial_auto_below_min:{result.initial_auto_fix_count}"
            )

    if "final_auto" in expected:
        if result.final_auto_fix_count != expected["final_auto"]:
            failures.append(
                f"final_auto:{result.final_auto_fix_count}"
            )

    applied = _applied_count(result)
    if "min_applied" in expected and applied < expected["min_applied"]:
        failures.append(f"applied_below_min:{applied}")

    if "output_same" in expected:
        actual = result.output == source
        if actual != expected["output_same"]:
            failures.append(f"output_same:{actual}")

    if "final_contains" in expected:
        if expected["final_contains"] not in _all_text(result.output):
            failures.append("final_text_missing")

    if "min_final_suggestions" in expected:
        if (
            result.final_suggestion_count
            < expected["min_final_suggestions"]
        ):
            failures.append(
                f"final_suggestions_below_min:"
                f"{result.final_suggestion_count}"
            )

    if "new_suggestions" in expected:
        if len(result.new_suggestion_keys) != expected["new_suggestions"]:
            failures.append(
                f"new_suggestions:{len(result.new_suggestion_keys)}"
            )

    if "new_high" in expected:
        if len(result.new_high_impact_keys) != expected["new_high"]:
            failures.append(
                f"new_high:{len(result.new_high_impact_keys)}"
            )

    if "heading_contains" in expected:
        document = Document(BytesIO(result.output))
        headings = [
            paragraph.text
            for paragraph in document.paragraphs
            if (paragraph.style.name or "").lower().startswith("heading")
        ]
        if not any(
            expected["heading_contains"] in text
            for text in headings
        ):
            failures.append("heading_fix_missing")

    if "table_contains" in expected:
        document = Document(BytesIO(result.output))
        table_text = "\n".join(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        if expected["table_contains"] not in table_text:
            failures.append("table_fix_missing")

    if expected.get("bold_preserved"):
        document = Document(BytesIO(result.output))
        bold_fixed = any(
            "هذا" in run.text and run.bold is True
            for paragraph in document.paragraphs
            for run in paragraph.runs
        )
        if not bold_fixed:
            failures.append("bold_not_preserved")

    if "preserve_issue" in expected:
        issue_types = _semantic_issue_types(result.output)
        if expected["preserve_issue"] not in issue_types:
            failures.append(
                f"semantic_issue_not_preserved:{sorted(issue_types)}"
            )

    if expected.get("preserve_fact_conflict"):
        if _fact_conflict_count(result.output) < 1:
            failures.append("fact_conflict_not_preserved")

    if (
        expected.get("second_output_same")
        or "second_initial_auto" in expected
        or "second_applied" in expected
    ):
        second = run_docx_review_loop(result.output)
        if expected.get("second_output_same"):
            if second.output != result.output:
                failures.append("second_run_changed_bytes")
        if "second_initial_auto" in expected:
            if (
                second.initial_auto_fix_count
                != expected["second_initial_auto"]
            ):
                failures.append(
                    f"second_initial_auto:"
                    f"{second.initial_auto_fix_count}"
                )
        if "second_applied" in expected:
            second_applied = _applied_count(second)
            if second_applied != expected["second_applied"]:
                failures.append(
                    f"second_applied:{second_applied}"
                )
        if not second.stable or not second.safety_pass:
            failures.append("second_run_not_stable_safe")

    if result.new_suggestion_keys:
        failures.append(
            f"loop_created_new_suggestions:"
            f"{len(result.new_suggestion_keys)}"
        )
    if result.new_high_impact_keys:
        failures.append(
            f"loop_created_new_high_impact:"
            f"{len(result.new_high_impact_keys)}"
        )
    if not result.structure_preserved:
        failures.append("structure_not_preserved")

    return {
        "id": case["id"],
        "scenario": scenario,
        "passed": not failures,
        "failures": failures,
        "stable": result.stable,
        "safety_pass": result.safety_pass,
        "initial_auto_fix_count": result.initial_auto_fix_count,
        "final_auto_fix_count": result.final_auto_fix_count,
        "applied_count": applied,
        "rounds": len(result.rounds),
        "final_suggestion_count": result.final_suggestion_count,
        "new_suggestions": len(result.new_suggestion_keys),
        "new_high_impact": len(result.new_high_impact_keys),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--suite",
        default="benchmark/review_loop/v1.json",
    )
    parser.add_argument("--report")
    parser.add_argument("--enforce", action="store_true")
    args = parser.parse_args()

    payload = json.loads(
        Path(args.suite).read_text(encoding="utf-8")
    )
    results = [evaluate(case) for case in payload["cases"]]
    passed = sum(item["passed"] for item in results)
    total = len(results)
    score = passed / total if total else 0.0

    report = {
        "schema_version": 1,
        "suite": args.suite,
        "passed": passed,
        "total": total,
        "score": round(score, 4),
        "failures": [
            item for item in results if not item["passed"]
        ],
        "results": results,
    }

    print(
        f"Review Loop & Idempotence V1: "
        f"{passed}/{total} = {score:.1%}"
    )
    if report["failures"]:
        print("\nFailures:")
        for item in report["failures"]:
            print(
                f"- {item['id']}: "
                + ", ".join(item["failures"])
            )

    if args.report:
        Path(args.report).write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    if args.enforce and report["failures"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
