from __future__ import annotations

import argparse
import os
import statistics
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from io import BytesIO

import httpx
from docx import Document


def build_docx(paragraphs: int, words_per_paragraph: int) -> bytes:
    document = Document()
    document.add_heading("اختبار ضغط نَضِيد", level=1)

    vocabulary = [
        "المستند",
        "العربي",
        "المراجعة",
        "السياق",
        "المصطلحات",
        "الحقائق",
        "الصياغة",
        "التحرير",
        "المعنى",
        "الاتساق",
        "التحليل",
        "الأداء",
    ]

    for index in range(paragraphs):
        words = [
            vocabulary[(index + word_index) % len(vocabulary)]
            for word_index in range(words_per_paragraph)
        ]
        document.add_paragraph(" ".join(words) + ".")

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def percentile(values: list[float], p: float) -> float:
    if not values:
        return 0.0

    ordered = sorted(values)
    index = max(
        0,
        min(
            len(ordered) - 1,
            round((len(ordered) - 1) * p),
        ),
    )
    return ordered[index]


def run_one(
    url: str,
    token: str,
    payload: bytes,
    timeout_seconds: float,
) -> tuple[int, float, str]:
    started = time.perf_counter()

    try:
        response = httpx.post(
            url.rstrip("/") + "/v1/analyze/docx",
            headers={
                "Authorization": f"Bearer {token}",
            },
            files={
                "file": (
                    "load-test.docx",
                    payload,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            timeout=timeout_seconds,
        )
        elapsed = time.perf_counter() - started
        return response.status_code, elapsed, response.text[:160]
    except Exception as exc:
        elapsed = time.perf_counter() - started
        return 0, elapsed, str(exc)[:160]


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Concurrent load probe for the Nadid AEE service."
    )
    parser.add_argument(
        "--url",
        default=os.environ.get("AEE_BACKEND_URL"),
        help="AEE base URL; defaults to AEE_BACKEND_URL.",
    )
    parser.add_argument(
        "--token",
        default=os.environ.get("AEE_INTERNAL_TOKEN"),
        help="AEE token; defaults to AEE_INTERNAL_TOKEN.",
    )
    parser.add_argument("--requests", type=int, default=20)
    parser.add_argument("--concurrency", type=int, default=5)
    parser.add_argument("--paragraphs", type=int, default=250)
    parser.add_argument("--words-per-paragraph", type=int, default=70)
    parser.add_argument("--timeout", type=float, default=180.0)
    parser.add_argument(
        "--max-error-rate",
        type=float,
        default=0.01,
        help="Maximum tolerated failed-request ratio.",
    )
    parser.add_argument(
        "--max-p95",
        type=float,
        default=60.0,
        help="Maximum tolerated p95 latency in seconds.",
    )
    args = parser.parse_args()

    if not args.url:
        parser.error("--url or AEE_BACKEND_URL is required")

    if not args.token:
        parser.error("--token or AEE_INTERNAL_TOKEN is required")

    if args.requests < 1 or args.concurrency < 1:
        parser.error("requests and concurrency must be positive")

    payload = build_docx(
        paragraphs=args.paragraphs,
        words_per_paragraph=args.words_per_paragraph,
    )

    print(
        f"payload={len(payload) / 1024 / 1024:.2f}MB "
        f"requests={args.requests} "
        f"concurrency={args.concurrency}"
    )

    results: list[tuple[int, float, str]] = []

    with ThreadPoolExecutor(max_workers=args.concurrency) as pool:
        futures = [
            pool.submit(
                run_one,
                args.url,
                args.token,
                payload,
                args.timeout,
            )
            for _ in range(args.requests)
        ]

        for future in as_completed(futures):
            results.append(future.result())

    durations = [elapsed for _, elapsed, _ in results]
    successes = [item for item in results if item[0] == 200]
    failures = [item for item in results if item[0] != 200]

    error_rate = len(failures) / len(results)

    print(f"success={len(successes)}/{len(results)}")
    print(f"error_rate={error_rate:.2%}")
    print(f"p50={percentile(durations, 0.50):.3f}s")
    print(f"p95={percentile(durations, 0.95):.3f}s")
    print(f"p99={percentile(durations, 0.99):.3f}s")
    print(f"mean={statistics.mean(durations):.3f}s")

    if failures:
        print("sample_failures:")
        for status, elapsed, detail in failures[:5]:
            print(
                f"  status={status} elapsed={elapsed:.3f}s "
                f"detail={detail!r}"
            )

    if error_rate > args.max_error_rate:
        print("FAIL: error rate exceeds threshold")
        return 1

    if percentile(durations, 0.95) > args.max_p95:
        print("FAIL: p95 exceeds threshold")
        return 1

    print("PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
